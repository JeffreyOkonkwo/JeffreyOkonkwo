"""Delivery_Health_Model.xlsx: sprint scorecard with live formulas, a team summary, flow data and a capacity planner."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
OUT = '../excel/Delivery_Health_Model.xlsx'
SC = pd.read_csv('../powerbi/fact_sprint.csv'); FL = pd.read_csv('../powerbi/fact_flow.csv')
INK, ACC, LINE = '0C0E0D', 'FF4A1C', 'D3D6D1'
F = lambda **k: Font(name='Arial', size=k.pop('size', 10), **k)
HDR = PatternFill('solid', fgColor=INK); BAND = PatternFill('solid', fgColor='F3F4F1'); INP = Font(name='Arial', size=10, color='0000FF')
thin = Border(bottom=Side(style='thin', color=LINE))
wb = Workbook()

def header(ws, row, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row, j, c); x.font = F(bold=True, color='FFFFFF'); x.fill = HDR; x.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 30

# --- Read me
ws = wb.active; ws.title = 'Read me'
rows = [('Delivery Health Model', F(bold=True, size=16)),
        ('Three WiredTiger storage engine teams, two-week sprints, April 2025 to October 2026. Source: MongoDB public Jira, project WT.', F()),
        ('', F()),
        ('Sheets', F(bold=True)),
        ('Summary: say-do ratio, scope added and carry-over by team, all live formulas over the Sprints sheet.', F()),
        ('Sprints: one row per sprint. Blue cells are data, black cells are formulas.', F()),
        ('Flow: one row per finished item, with cycle time split by status.', F()),
        ('Capacity planner: the formula from the Sprint Planning Field Guide. Change the blue inputs.', F()),
        ('', F()),
        ('Definitions', F(bold=True)),
        ('Committed: in the sprint two hours after it started, rebuilt from each ticket history.', F()),
        ('Say-do ratio: committed items done by sprint end, divided by committed items.', F()),
        ('Scope added: items that joined the sprint after it started, divided by committed items.', F()),
        ('Built by Jeffrey Okonkwo.', F(italic=True))]
for i, (t, f) in enumerate(rows, 1): ws.cell(i, 1, t).font = f
ws.column_dimensions['A'].width = 110

# --- Sprints
sp = wb.create_sheet('Sprints')
cols = ['Team', 'Sprint', 'Start', 'End', 'Half', 'Committed', 'Committed done', 'Added', 'Added done', 'Carried over', 'Removed', 'Say-do ratio', 'Scope added %', 'Carry-over %', 'Hit 80%?']
header(sp, 1, cols)
for i, r in enumerate(SC.sort_values(['team', 'start']).itertuples(), 2):
    vals = [r.team, r.sprint, r.start, r.end, r.half, r.committed, r.committed_done, r.added, r.added_done, r.carried_over, r.removed]
    for j, v in enumerate(vals, 1):
        c = sp.cell(i, j, v); c.font = INP if j >= 6 else F(); c.border = thin
    sp.cell(i, 12, f'=IF(F{i}=0,"",G{i}/F{i})').number_format = '0%'
    sp.cell(i, 13, f'=IF(F{i}=0,"",H{i}/F{i})').number_format = '0%'
    sp.cell(i, 14, f'=IF(F{i}+H{i}=0,"",J{i}/(F{i}+H{i}))').number_format = '0%'
    sp.cell(i, 15, f'=IF(L{i}="","",IF(L{i}>=0.8,"Yes","No"))')
    for j in range(12, 16): sp.cell(i, j).font = F(); sp.cell(i, j).border = thin
n = len(SC) + 1
sp.conditional_formatting.add(f'L2:L{n}', CellIsRule(operator='lessThan', formula=['0.5'], font=Font(color=ACC, bold=True)))
for j, w in enumerate([16, 30, 11, 11, 16, 11, 11, 9, 10, 11, 9, 11, 11, 11, 9], 1): sp.column_dimensions[L(j)].width = w
sp.freeze_panes = 'C2'; sp.auto_filter.ref = f'A1:O{n}'

# --- Summary
sm = wb.create_sheet('Summary', 1)
sm['A1'] = 'Delivery health by team'; sm['A1'].font = F(bold=True, size=14)
header(sm, 3, ['Team', 'Sprints', 'Committed', 'Committed done', 'Say-do ratio', 'Added', 'Scope added %', 'Carried over', 'Sprints at 80%+'])
teams = ['SE Foundations', 'SE Persistence', 'SE Transactions']
for i, t in enumerate(teams, 4):
    sm.cell(i, 1, t).font = F(bold=True)
    sm.cell(i, 2, f'=COUNTIF(Sprints!$A$2:$A${n},A{i})')
    sm.cell(i, 3, f'=SUMIF(Sprints!$A$2:$A${n},A{i},Sprints!$F$2:$F${n})')
    sm.cell(i, 4, f'=SUMIF(Sprints!$A$2:$A${n},A{i},Sprints!$G$2:$G${n})')
    sm.cell(i, 5, f'=IF(C{i}=0,"",D{i}/C{i})').number_format = '0.0%'
    sm.cell(i, 6, f'=SUMIF(Sprints!$A$2:$A${n},A{i},Sprints!$H$2:$H${n})')
    sm.cell(i, 7, f'=IF(C{i}=0,"",F{i}/C{i})').number_format = '0.0%'
    sm.cell(i, 8, f'=SUMIF(Sprints!$A$2:$A${n},A{i},Sprints!$J$2:$J${n})')
    sm.cell(i, 9, f'=COUNTIFS(Sprints!$A$2:$A${n},A{i},Sprints!$O$2:$O${n},"Yes")')
r = 7; sm.cell(r, 1, 'Program').font = F(bold=True)
for j, c in [(2, 'B'), (3, 'C'), (4, 'D'), (6, 'F'), (8, 'H'), (9, 'I')]: sm.cell(r, j, f'=SUM({c}4:{c}6)').font = F(bold=True)
sm.cell(r, 5, '=D7/C7').number_format = '0.0%'; sm.cell(r, 7, '=F7/C7').number_format = '0.0%'
for j in (5, 7): sm.cell(r, j).font = F(bold=True, color=ACC)
sm['A10'] = 'Say-do ratio by half year'; sm['A10'].font = F(bold=True, size=12)
header(sm, 11, ['Team', 'Apr to Dec 2025', 'Jan to Oct 2026', 'Change'])
for i, t in enumerate(teams, 12):
    sm.cell(i, 1, t)
    for j, h in [(2, 'Apr to Dec 2025'), (3, 'Jan to Oct 2026')]:
        sm.cell(i, j, f'=IFERROR(SUMIFS(Sprints!$G$2:$G${n},Sprints!$A$2:$A${n},$A{i},Sprints!$E$2:$E${n},"{h}")/SUMIFS(Sprints!$F$2:$F${n},Sprints!$A$2:$A${n},$A{i},Sprints!$E$2:$E${n},"{h}"),"")').number_format = '0.0%'
    sm.cell(i, 4, f'=IFERROR(C{i}-B{i},"")').number_format = '+0.0%;-0.0%'
for j, w in enumerate([18, 14, 14, 15, 13, 10, 14, 13, 14], 1): sm.column_dimensions[L(j)].width = w

# --- Flow
fw = wb.create_sheet('Flow')
header(fw, 1, ['Item', 'Team', 'Type', 'Started', 'Done', 'Cycle days', 'Working days', 'Code review days', 'Blocked days', 'Other waiting days', 'Flow efficiency'])
for i, r in enumerate(FL.itertuples(), 2):
    for j, v in enumerate([r.key, r.team, r.type, r.started, r.done, r.cycle_days, r.in_progress_days, r.code_review_days, r.blocked_days, r.other_wait_days], 1):
        fw.cell(i, j, v).font = INP if j >= 6 else F()
    fw.cell(i, 11, f'=IF(F{i}=0,"",G{i}/F{i})').number_format = '0%'
m = len(FL) + 1
fw.cell(m + 2, 5, 'Median').font = F(bold=True); fw.cell(m + 2, 6, f'=MEDIAN(F2:F{m})').number_format = '0.0'
fw.cell(m + 3, 5, '85th pct').font = F(bold=True); fw.cell(m + 3, 6, f'=PERCENTILE(F2:F{m},0.85)').number_format = '0.0'
fw.cell(m + 4, 5, 'Review share').font = F(bold=True); fw.cell(m + 4, 6, f'=SUM(H2:H{m})/SUM(F2:F{m})').number_format = '0.0%'
for j, w in enumerate([11, 16, 14, 11, 11, 10, 12, 15, 12, 17, 14], 1): fw.column_dimensions[L(j)].width = w
fw.freeze_panes = 'B2'

# --- Capacity planner
cp = wb.create_sheet('Capacity planner')
cp['A1'] = 'Sprint capacity planner'; cp['A1'].font = F(bold=True, size=14)
cp['A2'] = 'From the Sprint Planning Field Guide. Change the blue cells.'; cp['A2'].font = F(italic=True)
inp = [('People on the team', 6), ('Working days in the sprint', 10), ('Leave and holiday days (team total)', 4), ('Meeting days (team total)', 6), ('Focus factor', 0.7), ('Commit level', 0.8), ('Unplanned work buffer', 0.2)]
for i, (k, v) in enumerate(inp, 4):
    cp.cell(i, 1, k).font = F(); c = cp.cell(i, 2, v); c.font = INP
    if isinstance(v, float): c.number_format = '0%'
out = [('Raw person-days', '=B4*B5'), ('Available person-days', '=B11-B6-B7'), ('Real capacity (after focus factor)', '=B12*B8'),
       ('Reserved for unplanned work', '=B13*B10'), ('Commit up to (person-days)', '=(B13-B14)*B9')]
for i, (k, f) in enumerate(out, 11):
    cp.cell(i, 1, k).font = F(bold=True); c = cp.cell(i, 2, f); c.number_format = '0.0'; c.font = F(bold=True, color=ACC if i == 15 else INK)
cp.column_dimensions['A'].width = 40; cp.column_dimensions['B'].width = 14
wb.save(OUT); print('saved', OUT)
