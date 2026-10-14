"""Build Will_My_Vote_Count.xlsx: live formulas over the tagged comments, the 2023 result sheets and the fixes."""
import os, sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
sys.path.insert(0, './scripts')
from themes import THEMES

NG = './'
OUT = os.environ.get('OUT', '/mnt/user-data/outputs/Will_My_Vote_Count.xlsx')
NAVY, CREAM, MUST, GREEN, RED, GREY, SOFT = '1E2A47', 'F4EEDF', 'E1B13B', '2F6B4F', 'D7322B', 'C9CDD6', '5B6478'
F = 'Arial'
fill = lambda c: PatternFill('solid', fgColor=c)
thin = Side(style='thin', color='E2DCCB')

def head(ws, row, cols, start=1):
    for i, h in enumerate(cols):
        c = ws.cell(row, start + i, h)
        c.font = Font(F, bold=True, color='FFFFFF', size=10); c.fill = fill(NAVY)
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 30

def body(c, bold=False, color='1B1F2A', size=10, wrap=False, fmt=None, align=None):
    c.font = Font(F, bold=bold, color=color, size=size)
    c.alignment = Alignment(vertical='top', wrap_text=wrap, horizontal=align)
    c.border = Border(bottom=thin)
    if fmt: c.number_format = fmt

def title(ws, text, sub):
    ws['A1'] = text; ws['A1'].font = Font(F, bold=True, size=18, color=NAVY)
    ws['A2'] = sub; ws['A2'].font = Font(F, size=10, color=SOFT)
    ws.sheet_view.showGridLines = False

ids = [t[0] for t in THEMES]
fears = pd.read_csv(NG + 'dash/data/fears.csv')
themes = pd.read_csv(NG + 'dash/data/themes.csv')
steps = pd.read_csv(NG + 'dash/data/steps.csv')
res = pd.read_csv(NG + 'dash/data/results_2023.csv')
fixes = pd.read_csv(NG + 'dash/data/fixes.csv')
ctx = pd.read_csv(NG + 'dash/data/context.csv')
videos = pd.read_csv(NG + 'dash/data/videos.csv')

wb = Workbook()

# ---------- Fears data (one row per fear comment) ----------
fd = wb.active; fd.title = 'Fears data'
title(fd, 'Fear comments, tagged', 'One row per public YouTube comment that names a fear (April 2025 to October 2026). 1 = the comment names that theme. Text and author names are not stored.')
cols = ['id', 'month', 'channel', 'video_id', 'tone'] + ids
step_cols = [f'step_{s.split()[0]}' for s in steps.step]
head(fd, 4, cols + step_cols)
for r, row in enumerate(fears.itertuples(index=False), start=5):
    for c, v in enumerate(row, start=1):
        fd.cell(r, c, v).font = Font(F, size=9)
    # step helper: 1 when any theme of that step is flagged
    for k, s in enumerate(steps.step):
        tids = themes[themes.process_step == s].theme.tolist()
        refs = ','.join(f'{L(6 + ids.index(t))}{r}=1' for t in tids)
        fd.cell(r, len(cols) + 1 + k, f'=IF(OR({refs}),1,0)' if tids else 0).font = Font(F, size=9, color=SOFT)
N_LAST = 4 + len(fears)
fd.freeze_panes = 'B5'
for i in range(1, len(cols) + len(step_cols) + 1): fd.column_dimensions[L(i)].width = 11
fd.column_dimensions['A'].width = 14; fd.column_dimensions['C'].width = 22
rng = lambda col: f"'Fears data'!${col}$5:${col}${N_LAST}"
col_of = {t: L(6 + i) for i, t in enumerate(ids)}
step_col = {s: L(len(cols) + 1 + k) for k, s in enumerate(steps.step)}

# ---------- Themes ----------
th = wb.create_sheet('Themes')
title(th, 'Rigging at collation is the fear people raise everywhere', 'Share of fear comments naming each theme, against the 2023 results week on X. Formulas count the Fears data sheet.')
head(th, 4, ['Theme', 'INEC step', 'Comments', 'Share now', 'Videos with theme', '2023 share (X)', 'Change (points)'])
for r, t in enumerate(themes.itertuples(), start=5):
    body(th.cell(r, 1, t.label), bold=True); body(th.cell(r, 2, t.process_step), color=SOFT)
    body(th.cell(r, 3, f'=COUNTIF({rng(col_of[t.theme])},1)'), fmt='#,##0')
    body(th.cell(r, 4, f'=C{r}/COUNTA({rng("A")})'), fmt='0.0%')
    vids = fears.loc[fears[t.theme] == 1, 'video_id'].nunique()
    body(th.cell(r, 5, vids), fmt='0')
    if pd.isna(t.y2023_share):
        body(th.cell(r, 6, 'n/a'), color=SOFT, align='right'); body(th.cell(r, 7, 'n/a'), color=SOFT, align='right')
    else:
        body(th.cell(r, 6, float(t.y2023_share)), fmt='0.0%', color='0000FF')
        body(th.cell(r, 7, f'=(D{r}-F{r})*100'), fmt='+0.0;-0.0;0.0')
TL = 4 + len(themes)
th.cell(TL + 1, 1, 'Blue cells are inputs from the 2023 baseline (10,000 posts on X, 26 Feb to 14 Mar 2023, Kaggle). "Videos with theme" was counted in Python from the same rows. INEC distrust has no 2023 figure because those posts were collected by searching for INEC.').font = Font(F, size=9, italic=True, color=SOFT)
th.conditional_formatting.add(f'D5:D{TL}', DataBarRule(start_type='num', start_value=0, end_type='max', color=MUST))
for c, w in zip('ABCDEFG', [38, 30, 11, 11, 15, 15, 14]): th.column_dimensions[c].width = w
ch = BarChart(); ch.type = 'bar'; ch.style = 10; ch.title = 'Share of fear comments by theme'
ch.add_data(Reference(th, min_col=4, min_row=4, max_row=TL), titles_from_data=True)
ch.set_categories(Reference(th, min_col=1, min_row=5, max_row=TL))
ch.y_axis.numFmt = '0%'; ch.y_axis.majorGridlines = None; ch.x_axis.scaling.orientation = 'maxMin'
ch.series[0].graphicalProperties.solidFill = MUST; ch.series[0].graphicalProperties.line.noFill = True
ch.legend = None; ch.height = 10; ch.width = 18; ch.gapWidth = 40
ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True; ch.dataLabels.showSerName = False; ch.dataLabels.showCatName = False; ch.dataLabels.showLegendKey = False; ch.dataLabels.numFmt = '0%'
th.add_chart(ch, f'A{TL + 3}')

# ---------- Journey (the dashboard sheet) ----------
j = wb.create_sheet('Journey', 0)
j.sheet_view.showGridLines = False
for c in 'ABCDEFGHIJ': j.column_dimensions[c].width = 15
j.column_dimensions['A'].width = 3
img = XLImage(NG + 'art/xl_banner.png'); img.width, img.height = 1180, 236; j.add_image(img, 'B1')
for r in range(1, 14): j.row_dimensions[r].height = 15
j['B14'] = 'Will My Vote Count?'; j['B14'].font = Font(F, bold=True, size=22, color=NAVY)
j['B15'] = 'Which steps in INEC\'s process cause the most public fear, and which fixes would rebuild trust before the 16 January 2027 vote?'; j['B15'].font = Font(F, size=11, color=SOFT)
kp = [('B', 'Comments read', f"=SUM(Videos!E5:E{4 + len(videos)})", '#,##0', NAVY),
      ('D', 'Raised a specific concern', f'=COUNTA({rng("A")})', '#,##0', NAVY),
      ('F', 'About upload, collation or courts', f'=SUMPRODUCT(--((({rng(step_col[steps.step[4]])})+({rng(step_col[steps.step[5]])})+({rng(step_col[steps.step[6]])}))>0))/D18', '0%', 'B07A10'),
      ('H', '2023 units with no usable result', "='2023 results'!G43", '0.0%', RED)]
for col, k, f, fmt, color in kp:
    j[f'{col}17'] = k; j[f'{col}17'].font = Font(F, size=9, color=SOFT)
    c = j[f'{col}18']; c.value = f; c.number_format = fmt; c.font = Font(F, bold=True, size=24, color=color)
j.row_dimensions[18].height = 34
j['B20'] = 'The journey of one vote'; j['B20'].font = Font(F, bold=True, size=14, color=NAVY)
j['B21'] = 'Share of fear comments touching each step. A comment can touch several steps.'; j['B21'].font = Font(F, size=9, color=SOFT)
head(j, 22, ['Step', 'What it is', 'Fear comments', 'Share', 'What fails'], start=2)
j.merge_cells('F22:J22')
for k, s in enumerate(steps.itertuples(), start=23):
    body(j.cell(k, 2, int(s.order)), bold=True, align='center'); body(j.cell(k, 3, s.short), bold=True)
    body(j.cell(k, 4, f'=SUM({rng(step_col[s.step])})'), fmt='0')
    body(j.cell(k, 5, f'=D{k}/$D$18'), fmt='0%', bold=True)
    body(j.cell(k, 6, s.what_fails), wrap=True); j.merge_cells(f'F{k}:J{k}')
    j.row_dimensions[k].height = 30
j.conditional_formatting.add('E23:E29', DataBarRule(start_type='num', start_value=0, end_type='max', color=MUST))
j.conditional_formatting.add('E23:E29', FormulaRule(formula=['E23=MAX($E$23:$E$29)'], font=Font(F, bold=True, color=RED)))
jc = BarChart(); jc.type = 'col'; jc.style = 10; jc.title = 'Fear by step'
jc.add_data(Reference(j, min_col=5, min_row=22, max_row=29), titles_from_data=True)
jc.set_categories(Reference(j, min_col=3, min_row=23, max_row=29))
jc.series[0].graphicalProperties.solidFill = MUST; jc.series[0].graphicalProperties.line.noFill = True
jc.legend = None; jc.y_axis.numFmt = '0%'; jc.y_axis.majorGridlines = None; jc.height = 7; jc.width = 16; jc.gapWidth = 50
jc.dataLabels = DataLabelList(); jc.dataLabels.showVal = True; jc.dataLabels.showSerName = False; jc.dataLabels.showCatName = False; jc.dataLabels.showLegendKey = False; jc.dataLabels.numFmt = '0%'
j.add_chart(jc, 'B31')
# step picker
j['G31'] = 'Pick a step'; j['G31'].font = Font(F, bold=True, size=12, color=NAVY)
j['G32'] = steps.short[5]; j['G32'].font = Font(F, bold=True, size=12, color='1B1F2A'); j['G32'].fill = fill('FFF6D8')
j['G32'].border = Border(top=Side('thin', color=MUST), bottom=Side('thin', color=MUST), left=Side('thin', color=MUST), right=Side('thin', color=MUST))
dv = DataValidation(type='list', formula1='=$C$23:$C$29', allow_blank=False); j.add_data_validation(dv); dv.add('G32')
j['H32'] = '< change me'; j['H32'].font = Font(F, size=9, italic=True, color=SOFT)
j['G34'] = 'Evidence'; j['G34'].font = Font(F, bold=True, size=10, color=SOFT)
j['G35'] = '=INDEX(Steps!E5:E11,MATCH(G32,Steps!C5:C11,0))'
j['G35'].font = Font(F, size=10); j['G35'].alignment = Alignment(wrap_text=True, vertical='top'); j.merge_cells('G35:J38')
j['G40'] = 'Fixes for this step'; j['G40'].font = Font(F, bold=True, size=10, color=SOFT)
for n in range(1, 4):
    r = 40 + n
    j[f'G{r}'] = ArrayFormula(f'G{r}', f'=IFERROR(INDEX(Fixes!$D$5:$D$19,SMALL(IF(Fixes!$C$5:$C$19=INDEX(Steps!$B$5:$B$11,MATCH($G$32,Steps!$C$5:$C$11,0)),ROW(Fixes!$C$5:$C$19)-4),{n})),"")')
    j[f'G{r}'].font = Font(F, size=10); j[f'G{r}'].alignment = Alignment(wrap_text=True, vertical='top'); j.merge_cells(f'G{r}:J{r}')
    j.row_dimensions[r].height = 42
j['B47'] = 'Read me: Amaka is an imaginary first-time voter in Enugu. Data: public YouTube comments (vidIQ), CCIJ audit of 2023 IReV result sheets, Yiaga Africa and Afrobarometer surveys. This is a reading of public fear, not a poll.'
j['B47'].font = Font(F, size=9, italic=True, color=SOFT)

# ---------- Steps ----------
st = wb.create_sheet('Steps')
title(st, 'INEC process steps', 'What fails at each step and the evidence. Sources: INEC, Yiaga Africa, EU EOM 2023, Dataphyte, press reports (see Sources).')
head(st, 4, ['Order', 'Step', 'Short name', 'What fails', 'Evidence'])
for r, s in enumerate(steps.itertuples(), start=5):
    for c, v in enumerate([s.order, s.step, s.short, s.what_fails, s.evidence], start=1): body(st.cell(r, c, v), wrap=True)
    st.row_dimensions[r].height = 44
for c, w in zip('ABCDE', [7, 28, 13, 40, 70]): st.column_dimensions[c].width = w

# ---------- 2023 results ----------
rs = wb.create_sheet('2023 results')
title(rs, 'In 2023, one polling unit in seven had no usable result online', 'Presidential result sheets on IReV by state. Source: Center for Collaborative Investigative Journalism audit of all polling units (GitHub).')
head(rs, 4, ['State', 'Polling units', 'Never uploaded', 'Unreadable', 'Wrong document', 'Problem units', 'Problem share', 'Missing signature or stamp'])
res = res.sort_values('state')
for r, s in enumerate(res.itertuples(), start=5):
    body(rs.cell(r, 1, s.state), bold=True)
    for c, v in zip(range(2, 6), [s.polling_units, s.never_uploaded, s.unreadable, s.wrong_document]): body(rs.cell(r, c, int(v)), fmt='#,##0', color='0000FF')
    body(rs.cell(r, 6, f'=SUM(C{r}:E{r})'), fmt='#,##0'); body(rs.cell(r, 7, f'=F{r}/B{r}'), fmt='0.0%', bold=True)
    body(rs.cell(r, 8, int(s.missing_signature_or_stamp)), fmt='#,##0', color='0000FF')
RL = 4 + len(res)
assert RL + 1 == 42
for c in range(2, 7):
    body(rs.cell(43, c, f'=SUM({L(c)}5:{L(c)}{RL})'), bold=True, fmt='#,##0')
body(rs.cell(43, 1, 'Nigeria'), bold=True); body(rs.cell(43, 7, '=F43/B43'), bold=True, fmt='0.0%', color=RED)
body(rs.cell(43, 8, f'=SUM(H5:H{RL})'), bold=True, fmt='#,##0')
rs.conditional_formatting.add(f'G5:G{RL}', DataBarRule(start_type='num', start_value=0, end_type='max', color='E9806F'))
rs.cell(45, 1, 'Blue cells are counts from the CCIJ dataset; black cells are formulas. "Problem" = never uploaded + unreadable + wrong document. Cancelled units are excluded from problems.').font = Font(F, size=9, italic=True, color=SOFT)
for c, w in zip('ABCDEFGH', [16, 13, 14, 12, 15, 13, 13, 16]): rs.column_dimensions[c].width = w
rs.freeze_panes = 'B5'

# ---------- Fixes ----------
fx = wb.create_sheet('Fixes')
title(fx, 'Most of the fixes need no new law and can land before 16 January', 'Each fix is tied to a fear theme and the step it changes. Filter the table with the header arrows.')
head(fx, 4, ['Fix', 'Theme', 'INEC step', 'What to do', 'Owner', 'Before January', 'Needs new law', 'Evidence'])
for r, f in enumerate(fixes.itertuples(), start=5):
    for c, v in enumerate([f.fix_id, f.theme, f.process_step, f.fix, f.owner, f.before_jan_2027, f.needs_new_law, f.evidence], start=1):
        body(fx.cell(r, c, v), wrap=True, bold=(c == 4))
    fx.row_dimensions[r].height = 46
fx.auto_filter.ref = f'A4:H{4 + len(fixes)}'
fx.conditional_formatting.add(f'F5:G{4 + len(fixes)}', FormulaRule(formula=['AND($F5="Yes",$G5="No")'], fill=fill('DCEFE3')))
fx['A21'] = 'Quick wins (before January, no new law)'; fx['A21'].font = Font(F, bold=True, color=GREEN)
fx['D21'] = f'=COUNTIFS(F5:F{4 + len(fixes)},"Yes",G5:G{4 + len(fixes)},"No")'; fx['D21'].font = Font(F, bold=True, size=14, color=GREEN)
for c, w in zip('ABCDEFGH', [6, 12, 24, 52, 20, 11, 11, 56]): fx.column_dimensions[c].width = w

# ---------- Videos ----------
vd = wb.create_sheet('Videos')
title(vd, 'Videos sampled', 'Every video whose comments were read. Comment index rows come from vidIQ, which hides channel names.')
head(vd, 4, ['Video ID', 'Channel', 'Title', 'Published', 'Comments read', 'Fear comments'])
for r, v in enumerate(videos.itertuples(), start=5):
    for c, val in enumerate([v.video_id, v.channel, v.title if isinstance(v.title, str) else '', v.published if isinstance(v.published, str) else '', int(v.comments_read)], start=1):
        body(vd.cell(r, c, val), fmt='#,##0' if c == 5 else None)
    body(vd.cell(r, 6, f'=COUNTIF({rng("D")},A{r})'), fmt='#,##0')
for c, w in zip('ABCDEF', [16, 30, 60, 12, 14, 14]): vd.column_dimensions[c].width = w

# ---------- Sources ----------
so = wb.create_sheet('Sources')
title(so, 'Sources and method', 'Every figure in this workbook traces to one of these.')
rows = [
    ('Survey figures', ''),
] + [(f'{c.metric}', f'{c.value} ({c.source}, {c.date})') for c in ctx.itertuples()] + [
    ('', ''),
    ('YouTube comments', '2,230 public comments under 70 videos (Channels Television, Arise News, TVC News, SaharaTV, four commentary channels, plus a comment index), collected through the vidIQ connector, April 2025 to October 2026. Author names were never stored.'),
    ('Theme tagging', 'Keyword rules in English and Pidgin (scripts/themes.py). Hand-checked sample: about 85 to 90 percent of tags correct. INEC distrust also counts comments naming INEC with a tone score of -0.3 or lower (VADER with a Nigerian lexicon).'),
    ('2023 baseline', '10,000 posts on X, 26 February to 14 March 2023 (Kaggle: twitter-data-nigerian-2023-presidential-election), anonymised.'),
    ('2023 result sheets', 'Center for Collaborative Investigative Journalism, nigeria-2023-election dataset (GitHub), all 176,846 polling units.'),
    ('Process evidence', 'INEC, Yiaga Africa, EU Election Observation Mission 2023, Situation Room, Dataphyte, ThisDay, Daily Trust, Channels Television; see research notes on GitHub.'),
    ('Limits', 'YouTube commenters are not representative of voters. English and Pidgin only. This measures fear, not voting intention. No party or candidate is scored.'),
]
for r, (a, b) in enumerate(rows, start=4):
    so.cell(r, 1, a).font = Font(F, bold=True, size=10); c = so.cell(r, 2, b); c.font = Font(F, size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
so.column_dimensions['A'].width = 40; so.column_dimensions['B'].width = 110

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {'Journey': MUST, 'Fixes': GREEN, '2023 results': RED}.get(ws.title, NAVY)
wb.active = 0
wb.save(OUT); print('saved', OUT)
